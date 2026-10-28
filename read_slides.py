from pptx import Presentation

prs = Presentation('Medical-Question-Answering-using-Llama-3-RAG.pptx')
for i, slide in enumerate(prs.slides):
    if i >= 6:
        break
    print(f'=== Slide {i+1} ===')
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                text = para.text.strip()
                if text:
                    print(f'  {repr(text)}')
